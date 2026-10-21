import base64
import datetime
import io
import json
import math
import re
import time
import urllib.parse
import uuid

import pandas as pd
import requests
import streamlit as st
from lxml import etree
from openai import APIStatusError, AuthenticationError, OpenAI, RateLimitError
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

st.set_page_config(page_title="医学文献 → 学术汇报 PPT", page_icon="🩺", layout="wide")

DEFAULT_MODEL = "nvidia/nemotron-3-super-120b-a12b:free"
FREE_MODELS = [  # 联网获取失败时的备用列表；平时会实时拉取 OpenRouter 当前可用的免费模型
    DEFAULT_MODEL,
    "openrouter/free",
    "google/gemma-4-31b-it:free",
]
PAID_MODELS = [
    "anthropic/claude-3.5-sonnet",
    "deepseek/deepseek-chat",
    "openai/gpt-4o",
    "google/gemini-pro-1.5",
]

# ---- 配色（参照参考 PPT）：浅灰底 + 深蓝 + 青蓝 + 珊瑚红 ----
NAVY = RGBColor(0x1E, 0x3A, 0x66)
CYAN = RGBColor(0x4F, 0xB5, 0xD9)
CORAL = RGBColor(0xE7, 0x6B, 0x6B)
GRAY = RGBColor(0x46, 0x56, 0x67)
INK = RGBColor(0x1F, 0x23, 0x29)
BG = RGBColor(0xF7, 0xF8, 0xFB)
LINE = RGBColor(0xDE, 0xE0, 0xE3)
CARD = RGBColor(0xE9, 0xED, 0xF3)
CORAL_TINT = RGBColor(0xF9, 0xE8, 0xEA)
ROW_ALT = RGBColor(0xEE, 0xF1, 0xF6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Microsoft YaHei"

# 版面（英寸，16:9）
SW, SH = 13.333, 7.5
X, W = 0.9, 11.53
Y0, Y1 = 2.15, 6.85

LAYOUTS = {
    "bullets": "要点列表",
    "abstract": "摘要（段落 + 关键词）",
    "pillars": "编号并列柱（01/02/03）",
    "chains": "链式通路（A→B→C）",
    "two_col": "左右分栏（对比）",
    "grid": "2×2 编号块",
    "cards": "卡片并列",
    "stats": "大数字指标",
    "flow": "流程步骤",
    "image": "原文图片 + 读数",
    "chart": "原生图表 + 读数",
    "table": "表格",
    "numbered": "编号结论",
}

BLOCK_LAYOUTS = ("two_col", "cards", "abstract", "pillars", "chains", "grid")

CAP_RE = re.compile(r"^\s*(fig(ure)?\.?\s*\d+|supplementary fig|图\s*\d+|附图)", re.I)
MIN_SIDE, MIN_AREA = 120, 30000


# ============================================================================
# 一、文献解析：文本 + 图片 + 表格
# ============================================================================
def rows_to_md(rows):
    return "\n".join("| " + " | ".join(r) + " |" for r in rows)


def norm_image(raw: bytes):
    """校验并统一为 PNG；过滤过小/极端长宽比的图（Logo、分隔线等）。"""
    try:
        im = Image.open(io.BytesIO(raw))
        im.load()
    except Exception:
        return None
    w, h = im.size
    if w < MIN_SIDE or h < MIN_SIDE or w * h < MIN_AREA or w / h > 8 or h / w > 8:
        return None
    if im.mode == "P":
        im = im.convert("RGBA")
    elif im.mode not in ("RGB", "RGBA", "L"):
        im = im.convert("RGB")
    im.thumbnail((1800, 1800))
    buf = io.BytesIO()
    im.save(buf, "PNG")
    return buf.getvalue(), im.size[0], im.size[1]


def find_caption(blocks, rect):
    best, bd = "", 1e9
    for b in blocks:
        if len(b) > 6 and b[6] != 0:
            continue
        x0, y0, x1, y1, txt = b[:5]
        txt = txt.strip().replace("\n", " ")
        if not CAP_RE.match(txt) or x1 < rect.x0 or x0 > rect.x1:
            continue
        if y0 >= rect.y1 - 20:
            d = y0 - rect.y1
        elif y1 <= rect.y0 + 20:
            d = (rect.y0 - y1) + 60
        else:
            continue
        if d < bd:
            best, bd = txt, d
    return best[:220] if bd < 260 else ""


def parse_pdf(data: bytes):
    import fitz  # PyMuPDF

    doc = fitz.open(stream=data, filetype="pdf")
    texts, imgs, seen = [], [], set()
    for pno, page in enumerate(doc, 1):
        texts.append(f"\n[第{pno}页]\n{page.get_text('text')}")
        try:
            for k, tb in enumerate(page.find_tables().tables):
                rows = [[(c or "").replace("\n", " ").strip() for c in r] for r in tb.extract()]
                if len(rows) >= 2 and len(rows[0]) >= 2:
                    texts.append(f"[原文表格 第{pno}页 #{k + 1}]\n" + rows_to_md(rows))
        except Exception:
            pass
        blocks = None
        for info in page.get_images(full=True):
            xref = info[0]
            if xref in seen:
                continue
            seen.add(xref)
            raw = None
            try:
                pix = fitz.Pixmap(doc, xref)
                if info[1]:
                    try:
                        pix = fitz.Pixmap(pix, fitz.Pixmap(doc, info[1]))
                    except Exception:
                        pass
                if pix.n - pix.alpha >= 4:
                    pix = fitz.Pixmap(fitz.csRGB, pix)
                raw = pix.tobytes("png")
            except Exception:
                try:
                    raw = doc.extract_image(xref)["image"]
                except Exception:
                    continue
            n = norm_image(raw)
            if not n:
                continue
            caption = ""
            try:
                rects = page.get_image_rects(xref)
                if rects:
                    if blocks is None:
                        blocks = page.get_text("blocks")
                    caption = find_caption(blocks, rects[0])
            except Exception:
                pass
            imgs.append(dict(id=len(imgs) + 1, page=pno, caption=caption, bytes=n[0], w=n[1], h=n[2]))
    return "\n".join(texts), imgs


def parse_docx(data: bytes):
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    d = docx.Document(io.BytesIO(data))
    out, imgs, pending = [], [], None
    for child in d.element.body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            txt = Paragraph(child, d).text.strip()
            for rid in child.xpath(".//a:blip/@r:embed"):
                try:
                    n = norm_image(d.part.related_parts[rid].blob)
                except Exception:
                    n = None
                if n:
                    imgs.append(dict(id=len(imgs) + 1, page=None, caption="", bytes=n[0], w=n[1], h=n[2]))
                    pending = imgs[-1]
                    out.append(f"[图片{pending['id']}位置]")
            if txt:
                if pending is not None and CAP_RE.match(txt):
                    pending["caption"] = txt[:220]
                pending = None
                out.append(txt)
        elif tag == "tbl":
            rows = [[c.text.strip().replace("\n", " ") for c in r.cells] for r in Table(child, d).rows]
            if rows:
                out.append("[原文表格]\n" + rows_to_md(rows))
    return "\n".join(out), imgs


def parse_text(data: bytes):
    for enc in ("utf-8", "gbk", "latin-1"):
        try:
            return data.decode(enc), []
        except Exception:
            continue
    return "", []


def parse_file(uploaded):
    data = uploaded.getvalue()
    ext = uploaded.name.rsplit(".", 1)[-1].lower()
    if ext == "pdf":
        return parse_pdf(data)
    if ext == "docx":
        return parse_docx(data)
    return parse_text(data)


def strip_references(text: str) -> str:
    ms = list(re.finditer(r"\n\s*(references|bibliography|参考文献)\s*\n", text, re.I))
    if ms and ms[-1].start() > len(text) * 0.5:
        return text[: ms[-1].start()]
    return text


# ============================================================================
# 二、大模型生成大纲（JSON）
# ============================================================================
SYSTEM_PROMPT = """你是一位资深的医学科研汇报助手，长期给科室组会和学术会议写 slides。你写的 PPT 被主任评价为“干净、有判断、不注水”。现在把一篇医学文献整理成完整的学术汇报内容。

【写作总原则——先读这一段，再动笔】
1. 一页一个核心论断：读者扫一眼标题就知道这页要论证什么；标题之后所有要点都在支撑这个论断，不堆砌无关信息。
2. 论断先行，证据紧随：每条要点写成“判断 + 关键证据/参数”的句式，不要先铺背景再给结论。
   - 反例（流水账，禁用）：“老年人压力感受器敏感性下降。”“麻醉药物会扩张血管。”“最终导致血压下降。”
   - 正例（有判断，推荐）：“压力感受器放电频率下降，使麻醉后反射性交感代偿失灵。”“舒张功能受限使 CO 对前负荷变化高度敏感——回心血量轻微下降即可触发低血压。”
3. 信息要增量：相邻两条要点不许同义反复；能合并的合并，能拆成更具体的就别写笼统话。
4. 禁用 AI 套话与空连接词：不要“综上所述/值得注意的是/不可否认/由此可见/首先其次最后/在当今社会/随着……的发展”；不要在每条开头加“关于/对于/在……方面”。
5. 书面语但不端着：写 slides 正文用简洁陈述句，少用“进行、开展、实施、相关、极具重要意义”这类大词；剂量、单位、阈值按医学惯例写全（0.05 µg·kg⁻¹·min⁻¹、MAP ≥ 65 mmHg、PPV > 10%）。
6. 区分事实与推断：文献直接报告的写事实；你整理出的因果关系用“提示/可能/提示与……相关”措辞，不要把综述性推测写成定论。

【第一步：判断文献类型，决定整体结构】
- 原始研究（RCT、队列、病例对照等）：背景与证据缺口 → 方法 → 结果（占多数）→ 讨论与意义 → 局限性 → 结论。
- 综述 / 指南 / 共识 / 知识更新：不要套用“方法-结果”结构，沿用文献自身章节顺序。典型顺序：摘要与关键词 → 背景与近年演进 → 机制 → 核心争议（正反方 + 未解决困境）→ 方法学瓶颈 → 临床转化建议（术前/诱导/维持/术后）→ 证据边界（已形成共识 vs 尚未确定）→ 未来方向。
- 其他类型按其内在逻辑组织。

【内容完整性】
1. 文献中每个有实质内容的章节都必须有对应幻灯片，不得遗漏（机制、方法学瓶颈、每一条争议、建议表格每一行、未来方向每一条）。页数不够时合并相近内容，绝不能丢弃整章。
2. 每页 3-6 个信息点，每点 1-2 句（≤60 字），写出文献里具体的机制、观点、参数、剂量、阈值、流程；不要只写“存在一定关联”“具有重要意义”这种空话。
3. 每页必须有标题和实际内容，不允许空页或只有标题。

【标题写法——最重要的一节】
标题必须是结论句或明确的研究问题，≤28 字，让读者不看正文也能抓住本页判断。
- 好标题：“老年 PIH 高发的三大生理归因”“争论一：术中 MAP 目标值如何设定？”“诱导前 5 min 启动去甲肾上腺素可减少血压波动”“四条通路共同决定术后 AKI 风险”。
- 坏标题（禁用）：“结果”“讨论”“机制部分”“关于低血压”“临床意义”“方法学”。
- 数字必须进标题时写完整：“MAP < 65 mmHg 持续 > 1 min 即触发再评估”。

【sections 的 subtitle】
每个章节的 subtitle 用一句话点出“本章要回答什么问题”，不要重复标题。
- 例：title="背景与生理基础"，subtitle="为什么老年患者在麻醉诱导后更容易低血压？"

【数据与证据】
4. 内容只能来自文献，严禁编造数据、试验名、文献、剂量、阈值或建议；文献没有的不要补。允许做解释性整理（机制写成因果链、争议整理成正反方对比）。
5. 有数据时写具体数值（n、%、均值±SD、中位数(IQR)、HR/OR/RR/MD、95%CI、P 值），与原文逐一一致。
6. 没有量化数据时，直接写文献里的机制、观点、参数和建议；不要在每条要点后重复“原文未报告”。证据局限只在【摘要页、证据边界页、结论页】集中说明一次，其余页最多用 callout 一句话提示（如“以下为专家共识，证据等级 C”）。文献标注了证据等级/推荐强度的建议，要在对应页面标出。

【版式 layout】（按内容选，不要整篇都用 bullets；相邻页尽量换版式）
- abstract：摘要页。blocks 为带标签段落（heading=标签，如“临床问题/研究焦点/证据边界”，bullets=1-2 句话）；items 为关键词（中英文）。
- pillars：2-4 个并列要点，带 01/02/03 编号。blocks：heading + 1-3 句论断。
- chains：多条因果链。blocks 每项一条链：heading=链名（如“通路 A · 神经反射”），bullets=按顺序的步骤（3-5 步，每步 ≤14 字），最后一步是结局。
- two_col：二元对比（正方/反方、已知/未知、已共识/尚未确定）。blocks 为 2 栏；争议类 contrast=true，“未解决的困境”写进 callout。
- grid：2×2 编号块。blocks 3-4 项：heading + 1-2 句论断。
- cards：3-4 个并列卡片（无编号）。blocks：heading + 1-3 句论断。
- stats：2-4 个关键数字或参数。stats 每项 {label, value 简短（如“>10%”“3.16 → 2.16”）, sub 用途/变化幅度}；可配 bullets 或 callout。
- flow：流程/决策路径/技术路线，3-6 步，steps {title, text}，最后一步为产出。**触发条件（必须用 flow，不要用 bullets）**：① 文献描述了"术前→诱导→维持→术后"这类时序操作链；② 出现"如果 MAP<65 则…否则…""PPV>10% 提示容量反应性"这类判断分支；③ 给出了诊断/治疗决策树、算法、评分流程。判断节点的 title 写成问句或条件（如"MAP < 65 mmHg？"），text 写分支动作；普通步骤 title 写动作。
- table：多环节/多指标对比。**触发条件（必须用 table，不要用 bullets 或 two_col）**：① 文献给了按"环节/时间点/分组"组织的建议表（如术前/诱导/维持/术后的措施+参数+证据等级）；② 原文本身就是表格（Table 1/2/3）；③ 要并列对比 ≥3 个分组/方案的多个维度。可重新加工：保留关键行列、合并或重命名表头；≤8 行 × 5 列，第一行表头，内容忠于原文。原文表格每行都要覆盖，行太多拆两页，每页都要写标题并带完整 table。**证据等级列**：如果是围术期建议表，最后一列写证据等级（A/B/C + 推荐强度），系统会自动上色。
- image：使用图片清单里的原文图（image_ids 引用），bullets 作为“关键读数”。
- chart：仅当文献明确给出 ≥3 个类别具体数值时使用，数值必须来自原文，严禁从图上目测估算。chart {type: bar/column/line, title, rows}，rows 第一行表头 ["类别","系列1",...]，其余每行 [类别, 数值, ...]。
- numbered：结论、未来方向等，3-5 条，用 items，每条一句话且本身是一个完整判断。
- bullets：常规要点（3-6 条）。

【其他字段】
7. sections 是目录章节（4-6 个，每个有 title 和一句话 subtitle；综述类沿用文献自身章节）；每页用 sec 指向章节序号（从 1 开始）。
8. callout 是页面底部一句话要点/提示（≤80 字），重要页面加，其余留空字符串。
9. notes 是口语化讲稿（3-6 句话），像站在台上对着主任讲：先承接上一页（“刚才讲了……，这一页我们看……”），再用一句大白话解释本页判断为什么重要，最后点一下证据出处（“见建议表格第 3 行”“数据来自 Table 2”）。不要照念 bullets，不要用书面长句。
10. 图片：只能用图片清单里的图，不要要求生成新图；无关的图不要用。
11. cover：kicker=文献类型 + 年份范围或期刊（如“深度综述 · 2021—2026”）；title=封面标题，可用 \\n 分行（每行≤14 字）；meta=一句话副标题（如“从机制、监测到闭环管理的全景梳理”）和作者单位（可多行）；foot=关键词或研究样本概括。
12. 只输出一个合法 JSON 对象，不要任何解释，不要 markdown 代码围栏；字符串中的换行用 \\n；所有字段值使用上面给出的名称和结构。

【自检——输出前在心里过一遍】
- 每个标题都是结论句，不是章节名？
- 每条 bullets 都有判断 + 具体参数/数字，不是流水账？
- 有没有出现“综上所述/值得注意的是/首先其次”这类套话？出现就删掉重写。
- 相邻页版式是否换过？是否有 1-2 页真的用了 chains/two_col/table/chart 而不是全部 bullets？
- notes 是否像人在说话，而不是正文复读？

【JSON 结构】（不用的字段可省略或给空值）
{
  "cover": {"kicker": "…", "title": "…", "meta": "…", "foot": "…"},
  "sections": [{"title": "背景与生理基础", "subtitle": "一句话点出本章要回答的问题"}],
  "slides": [
    {
      "sec": 1,
      "title": "结论式标题（≤28 字）",
      "layout": "two_col",
      "contrast": false,
      "sub_title": "",
      "bullets": ["判断 + 证据"],
      "blocks": [{"heading": "栏目/步骤组标题", "bullets": ["判断 + 证据"]}],
      "stats": [{"label": "指标名", "value": "3.16 → 2.16", "sub": "下降 32%"}],
      "steps": [{"title": "步骤名", "text": "说明"}],
      "items": ["完整判断句"],
      "image_ids": [1],
      "table": {"rows": [["环节", "措施", "证据等级"], ["…", "…", "…"]]},
      "chart": {"type": "bar", "title": "标题（含单位）", "rows": [["类别", "系列1"], ["A", "10"], ["B", "20"], ["C", "30"]]},
      "callout": "…",
      "notes": "口语化讲稿"
    }
  ],
  "references": ["作者. 题名. 期刊. 年;卷(期):页码."]
}"""


def extract_json(s: str):
    """宽松解析：去掉围栏/思考标签，失败时用 json-repair 修复（可处理缺逗号、被截断等）。"""
    s = re.sub(r"<think>.*?</think>", "", s, flags=re.S)
    s = re.sub(r"```(?:json)?", "", s)
    a = s.find("{")
    if a < 0:
        raise ValueError("模型没有返回 JSON")
    s = s[a:]
    for cand in (s[: s.rfind("}") + 1], re.sub(r",\s*([}\]])", r"\1", s[: s.rfind("}") + 1])):
        try:
            return json.loads(cand)
        except Exception:
            pass
    try:
        import json_repair

        obj = json_repair.loads(s)
        if isinstance(obj, dict) and obj.get("slides"):
            return obj
    except ImportError:
        pass
    raise ValueError("无法解析模型返回的 JSON")


@st.cache_data(ttl=86400, show_spinner=False)
def fetch_free_models():
    """实时获取 OpenRouter 当前的免费模型（按上下文长度从大到小），避免写死的模型 ID 过期。"""
    try:
        items = requests.get("https://openrouter.ai/api/v1/models", timeout=15).json()["data"]
        free = [(m["id"], m.get("context_length") or 0) for m in items if m["id"].endswith(":free")]
        free = [m for m in free if m[1] >= 32000]
        free.sort(key=lambda m: m[1], reverse=True)
        return [m[0] for m in free]
    except Exception:
        return []


# ---- 医学/麻醉学术语对照表（luovalley/ppt-web glossary.md）----
GLOSSARY_URL = "https://raw.githubusercontent.com/luovalley/ppt-web/main/glossary.md"


@st.cache_data(ttl=86400, show_spinner=False)
def load_glossary():
    """拉取并解析医学术语对照表：英文术语 -> 中文标准译法。
    格式：`* English Term — 中文译法`，容错 tab 分隔与备注行。"""
    try:
        raw = requests.get(GLOSSARY_URL, timeout=20).text
    except Exception:
        return {}
    gloss = {}
    for line in raw.splitlines():
        line = line.strip().lstrip("*").strip()
        if not line or line.startswith("#") or line.startswith("-") or line.startswith("|"):
            continue
        # 匹配 "English — 中文" 或 "English\t中文"
        m = re.match(r"^([A-Za-z][A-Za-z0-9\s\-\(\)/,]+?)\s*[—–\-]\s*([^\t]+?)\s*$", line)
        if not m:
            m = re.match(r"^([A-Za-z][A-Za-z0-9\s\-\(\)/,]+?)\t(.+?)\s*$", line)
        if not m:
            continue
        en = m.group(1).strip()
        zh = m.group(2).strip()
        # 只保留英文部分为纯医学术语（去掉括号内备注），译法保留第一个 "/" 前的主译
        if len(en) < 2 or len(zh) < 1:
            continue
        main_zh = zh.split("/")[0].strip()
        # 去掉译法里的括号注释
        main_zh = re.sub(r"\s*\([^)]*\)\s*", "", main_zh).strip()
        if main_zh:
            gloss[en.lower()] = main_zh
    return gloss


def pick_glossary_for_text(text, gloss, limit=120):
    """扫描文献正文，挑出在术语表里出现的英文术语，按出现频次返回对照表。"""
    if not gloss:
        return ""
    # 按术语长度从长到短排序，优先匹配多词术语（如 "Abdominal Compartment Syndrome"）
    terms = sorted(gloss.keys(), key=len, reverse=True)
    counts = {}
    # 用一个临时文本，逐术语计数
    lower = text.lower()
    for en in terms:
        # 用单词边界匹配，避免 "IVC" 命中 "IVCD"
        pattern = r"(?<![a-z])" + re.escape(en) + r"(?![a-z])"
        n = len(re.findall(pattern, lower))
        if n > 0:
            counts[en] = n
    if not counts:
        return ""
    # 按频次排序，取前 limit 个
    top = sorted(counts.items(), key=lambda x: -x[1])[:limit]
    lines = [f"{en} — {gloss[en]}" for en, _ in top]
    return "英文术语 — 中文标准译法（必须使用此译法，不要自创）：\n" + "\n".join(lines)


def api_error_text(e):
    """提取 OpenRouter 返回的真实错误信息（含上游供应商的原始报错）。"""
    msg = ""
    body = getattr(e, "body", None)
    if isinstance(body, dict):
        err = body.get("error", body)
        if isinstance(err, dict):
            msg = str(err.get("message", ""))
            raw = (err.get("metadata") or {}).get("raw")
            if raw:
                msg += f" | 上游：{str(raw)[:160]}"
            if err.get("metadata", {}).get("reasons"):
                msg += f" | 原因：{err['metadata']['reasons']}"
    return msg or str(e)


def call_llm(api_key, models, user_prompt, on_fail):
    """逐个模型尝试：限流(429)等待后重试；JSON 无效重试一次；模型不可用/出错则换下一个。"""
    client = OpenAI(
        api_key=api_key, base_url="https://openrouter.ai/api/v1", timeout=300, max_retries=0,
        default_headers={"HTTP-Referer": "https://streamlit.io", "X-Title": "Medical Literature to PPT"},
    )
    for m in models:
        for attempt in range(2):
            try:
                r = client.chat.completions.create(
                    model=m,
                    temperature=0.2,
                    messages=[
                        {"role": "system", "content": SYSTEM_PROMPT},
                        {"role": "user", "content": user_prompt},
                    ],
                )
                return extract_json(r.choices[0].message.content or ""), m
            except AuthenticationError:
                raise RuntimeError("API Key 无效或已失效，请检查。")
            except RateLimitError:
                if attempt == 0:
                    on_fail(m, "被限流(429)，8 秒后重试一次…")
                    time.sleep(8)
                    continue
                on_fail(m, "仍被限流，换下一个模型")
                break
            except ValueError as e:  # JSON 无效
                if attempt == 0:
                    on_fail(m, f"{e}，重试一次…")
                    continue
                on_fail(m, f"{e}，换下一个模型")
                break
            except APIStatusError as e:
                hint = {
                    403: "（常见原因：输入被内容审核拦截、账号隐私/数据策略限制、或服务器所在地区受限）",
                    404: "（模型已下架或与隐私设置不匹配，可到 openrouter.ai/settings/privacy 检查）",
                    402: "（额度不足）",
                }.get(e.status_code, "")
                on_fail(m, f"HTTP {e.status_code}：{api_error_text(e)[:260]}{hint}，换下一个模型")
                break
            except Exception as e:
                on_fail(m, f"{str(e)[:120]}，换下一个模型")
                break
    raise RuntimeError("所有候选模型都失败了：免费模型常被限流，请稍后再试，或在侧边栏换模型/填写自定义模型 ID。")


# ============================================================================
# 三、大纲数据结构与规范化
# ============================================================================
def to_int(x):
    m = re.search(r"\d+", str(x))
    return int(m.group()) if m else None


def as_list(x):
    if x is None or x == "":
        return []
    return x if isinstance(x, list) else [x]


def str_lines(v):
    if isinstance(v, str):
        v = v.split("\n")
    return [str(i).rstrip() for i in as_list(v) if str(i).strip()]


def norm_rows(t):
    rows = None
    if isinstance(t, dict):
        rows = t.get("rows")
        if rows and t.get("headers"):
            rows = [t["headers"]] + rows
    elif isinstance(t, list):
        rows = t
    if not rows:
        return None
    if isinstance(rows, list) and rows and isinstance(rows[0], dict):  # 兼容 [{列名: 值}, ...]
        keys = []
        for r in rows:
            keys += [k for k in r if k not in keys]
        rows = [keys] + [[r.get(k, "") for k in keys] for r in rows]
    rows = [[str(c) for c in r] for r in rows if isinstance(r, (list, tuple))]
    nc = max((len(r) for r in rows), default=0)
    if not nc:
        return None
    return [r + [""] * (nc - len(r)) for r in rows]


def norm_chart(c):
    if not isinstance(c, dict):
        return None, "bar", ""
    rows = c.get("rows")
    if not rows and c.get("categories") and c.get("series"):
        cats, ser = c["categories"], c["series"]
        rows = [[""] + [str(x.get("name", "")) for x in ser]]
        for i, cat in enumerate(cats):
            rows.append([str(cat)] + [str(x.get("values", [])[i]) if i < len(x.get("values", [])) else "" for x in ser])
    rows = norm_rows(rows)
    if not rows or len(rows) < 3 or len(rows[0]) < 2:
        return None, "bar", ""
    t = c.get("type", "bar")
    return rows, (t if t in ("bar", "column", "line") else "bar"), str(c.get("title", ""))


def new_slide(**kw):
    s = dict(
        uid=uuid.uuid4().hex[:8], sec=0, title="", layout="bullets", sub_title="", bullets=[],
        blocks=[], stats=[], steps=[], items=[], image_ids=[], contrast=False,
        table=None, table_cur=None, tv=0,
        chart=None, chart_cur=None, cv=0, chart_type="bar", chart_title="",
        callout="", notes="",
    )
    s.update(kw)
    return s


def parse_blocks_json(v):
    out = []
    for b in as_list(v):
        if isinstance(b, dict):
            out.append(dict(
                heading=str(b.get("heading") or b.get("title") or ""),
                lines=str_lines(b.get("bullets") or b.get("lines") or b.get("text") or b.get("steps")),
            ))
        elif isinstance(b, str) and b.strip():
            out.append(dict(heading="", lines=[b.strip()]))
    return out


def normalize_slide(x, images, nsec):
    bullets = str_lines(x.get("bullets"))
    blocks = parse_blocks_json(x.get("blocks"))
    stats = [
        dict(label=str(d.get("label", "")), value=str(d.get("value", "")), sub=str(d.get("sub", "")))
        for d in as_list(x.get("stats")) if isinstance(d, dict)
    ]
    steps = [
        dict(title=str(d.get("title", "")), text=str(d.get("text", "")))
        for d in as_list(x.get("steps")) if isinstance(d, dict)
    ]
    items = str_lines(x.get("items") or x.get("keywords"))
    ids = [to_int(i) for i in as_list(x.get("image_ids"))]
    ids = [i for i in ids if i in images]
    table = norm_rows(x.get("table"))
    chart, ctype, ctitle = norm_chart(x.get("chart"))

    # 没有任何实质内容的页面直接丢弃（避免出现空白页）
    if not (bullets or blocks or stats or steps or items or ids or table or chart):
        return None

    layout = x.get("layout")
    if layout not in LAYOUTS:
        layout = ("image" if ids else "table" if table else "chart" if chart else "two_col" if blocks
                  else "stats" if stats else "flow" if steps else "numbered" if items else "bullets")
    if ids and layout not in ("image", "table", "bullets"):
        if not bullets:
            bullets = [l for b in blocks for l in b["lines"]]
        layout = "image"
    need = {"table": table, "chart": chart, "stats": stats, "flow": steps, "numbered": items}
    need.update({k: blocks for k in BLOCK_LAYOUTS})
    if layout in need and not need[layout]:
        layout = "bullets"
    if layout == "bullets" and not bullets:
        bullets = [l for b in blocks for l in b["lines"]] or items or [d["text"] for d in steps]
    if layout == "bullets" and not bullets:
        return None

    title = str(x.get("title", "") or "").strip()
    if not title:  # 标题缺失时从内容推断
        title = (str(x.get("sub_title") or "").strip() or (blocks[0]["heading"] if blocks else "")
                 or (bullets[0][:24] if bullets else "") or "（请填写标题）")
    sec = to_int(x.get("sec")) or 0
    return new_slide(
        sec=sec if 0 <= sec <= nsec else 0, title=title, layout=layout,
        sub_title=str(x.get("sub_title", "") or ""), bullets=bullets, blocks=blocks, stats=stats, steps=steps,
        items=items, image_ids=ids, table=table, chart=chart, chart_type=ctype, chart_title=ctitle,
        contrast=bool(x.get("contrast", False)),
        callout=str(x.get("callout", "") or ""), notes=str(x.get("notes", "") or ""),
    )


def normalize_outline(data, images):
    cov = data.get("cover") or {}
    fix = lambda v: str(v or "").replace("\\n", "\n")
    cover = dict(
        kicker=fix(cov.get("kicker")), title=fix(cov.get("title") or data.get("title")),
        meta=fix(cov.get("meta") or data.get("subtitle")), foot=fix(cov.get("foot")),
    )
    secs = [
        dict(title=str(x.get("title", "")), subtitle=str(x.get("subtitle", "")))
        for x in as_list(data.get("sections")) if isinstance(x, dict)
    ]
    slides = [t for t in (normalize_slide(x, images, len(secs)) for x in as_list(data.get("slides")) if isinstance(x, dict)) if t]
    refs = [str(r) for r in as_list(data.get("references")) if str(r).strip()]
    return dict(
        uid=uuid.uuid4().hex[:8], cover=cover, sections=secs, slides=slides, references=refs,
        show_agenda=len(secs) >= 2, show_refs=bool(refs),
    )


# ---- 结构化文本 <-> 数据 ----
def blocks_to_text(blocks):
    return "\n".join(f"## {b['heading']}\n" + "\n".join(b["lines"]) for b in blocks)


def parse_blocks(text):
    blocks, cur = [], None
    for line in text.split("\n"):
        if line.startswith("##"):
            cur = dict(heading=line.lstrip("#").strip(), lines=[])
            blocks.append(cur)
        elif line.strip():
            if cur is None:
                cur = dict(heading="", lines=[])
                blocks.append(cur)
            cur["lines"].append(line.rstrip())
    return blocks


def split_bar(line):
    return [p.strip() for p in re.split(r"[|｜]", line)]


def stats_to_text(stats):
    return "\n".join(f"{d['label']} | {d['value']} | {d['sub']}" for d in stats)


def parse_stats(text):
    out = []
    for line in text.split("\n"):
        if line.strip():
            p = (split_bar(line) + ["", "", ""])[:3]
            out.append(dict(label=p[0], value=p[1], sub=p[2]))
    return out


def steps_to_text(steps):
    return "\n".join(f"{d['title']} | {d['text']}" for d in steps)


def parse_steps(text):
    out = []
    for line in text.split("\n"):
        if line.strip():
            p = (split_bar(line) + [""])[:2]
            out.append(dict(title=p[0], text=p[1]))
    return out


def final_rows(s, field):
    cur = "table_cur" if field == "table" else "chart_cur"
    rows = s.get(cur) or s.get(field)
    if not rows:
        return None
    rows = [r for r in rows if any(str(c).strip() for c in r)]
    return rows or None


def sec_label(o, sec):
    secs = o["sections"]
    return f"{sec:02d} · {secs[sec - 1]['title']}" if 1 <= sec <= len(secs) else ""


# ============================================================================
# 四、生成原生、可编辑的 PPTX（版式参照参考 PPT）
# ============================================================================
WIDE_PUNCT = set("–—→·≤≥±×≠“”‘’…～")


def text_w(s, size, bold=False):
    """估算文本宽度(pt)：中日韩字符与全角符号占 1 个字号宽，其余约 0.6（加粗 0.66）。"""
    a = 0.66 if bold else 0.6
    return sum(size if (ord(c) > 0x2E80 or c in WIDE_PUNCT) else size * a for c in s)


def text_height(lines, w_in, size, indent=36, after=0.5):
    avail = max(w_in * 72 - indent, 40)
    total = 0
    for text, lvl in lines:
        n = max(1, math.ceil(text_w(text, size) / max(avail - lvl * 22, 30)))
        total += n * size * 1.3 + size * after
    return total / 72


def parse_bullets(bullets):
    out = []
    for raw in bullets:
        lvl = 1 if raw.startswith(("\t", "  ")) else 0
        t = raw.strip().lstrip("-–•·*").strip()
        if t:
            out.append((t, lvl))
    return out


def style_run(run, size, bold=False, color=GRAY):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.color.rgb = color
    f.name = FONT
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", FONT)


def set_bullet(p, level):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(0.28 + 0.3 * level))))
    pPr.set("indent", str(-int(Inches(0.24))))
    clr = etree.SubElement(pPr, qn("a:buClr"))
    etree.SubElement(clr, qn("a:srgbClr")).set("val", "4FB5D9")
    etree.SubElement(pPr, qn("a:buChar")).set("char", "•" if level == 0 else "–")


def rect(slide, x, y, w, h, color):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    sh.line.fill.background()
    sh.shadow.inherit = False
    style = sh._element.find(qn("p:style"))  # 去掉主题样式引用，避免部分软件仍渲染阴影
    if style is not None:
        sh._element.remove(style)
    return sh


def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP, margin=0.04):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    return tf


def para(tf, k, text, size, bold=False, color=GRAY, align=PP_ALIGN.LEFT, after=0, bullet=None):
    p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
    p.alignment = align
    p.line_spacing = 1.1
    if after:
        p.space_after = Pt(after)
    r = p.add_run()
    r.text = text
    style_run(r, size, bold, color)
    if bullet is not None:
        set_bullet(p, bullet)
    return p


def add_text(slide, x, y, w, h, text, size, color=GRAY, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tf = textbox(slide, x, y, w, h, anchor)
    for k, line in enumerate(str(text).split("\n")):
        para(tf, k, line, size, bold, color, align)
    return tf


def set_title(ph, text, size, x, y, w, h, anchor, color=NAVY):
    ph.left, ph.top, ph.width, ph.height = Inches(x), Inches(y), Inches(w), Inches(h)
    tf = ph.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.text = text
    for p in tf.paragraphs:
        p.alignment = PP_ALIGN.LEFT
        for r in p.runs:
            style_run(r, size, True, color)


def add_bullets(slide, x, y, w, h, lines, color=GRAY, max_size=18):
    if not lines:
        return
    sizes = [z for z in (18, 16, 15, 14, 13, 12, 11, 10) if z <= max_size] or [10]
    size = next((z for z in sizes if text_height(lines, w, z) <= h - 0.1), sizes[-1])
    tf = textbox(slide, x, y, w, h)
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    for k, (t, lvl) in enumerate(lines):
        para(tf, k, t, size - 2 * lvl, False, color, after=size * 0.5, bullet=lvl)


def add_plain(slide, x, y, w, h, lines, color=GRAY, size=13):
    tf = textbox(slide, x, y, w, h)
    for k, t in enumerate(lines):
        para(tf, k, t, size, False, color, after=size * 0.4)


def add_panel(slide, x, y, w, h, sub_title, lines, max_size=16):
    """小标题（深蓝加粗）+ 要点列表。"""
    if h < 0.4:
        return
    if sub_title:
        add_text(slide, x, y, w, 0.4, sub_title, 15, NAVY, True)
        y, h = y + 0.5, h - 0.5
    if lines and h > 0.3:
        add_bullets(slide, x, y, w, h, lines, max_size=max_size)


def panel_need(sub_title, lines, w, size=15):
    return (0.5 if sub_title else 0) + (text_height(lines, w, size) if lines else 0)


def callout_h(text, w):
    n = max(1, math.ceil(text_w(text, 14) * 1.05 / max(w * 72 - 44, 60)))
    return min(1.7, max(0.8, n * 14 * 1.35 / 72 + 0.4))


def add_callout(slide, text, x, y, w, h):
    rect(slide, x, y, w, h, CORAL_TINT)
    add_text(slide, x + 0.28, y, w - 0.56, h, text, 14, NAVY, True, anchor=MSO_ANCHOR.MIDDLE)


NUM_RE = re.compile(r"^[<>≤≥~±+\-–]?\s*\d[\d.,%\s/()±\-–:]*$")


def add_table(slide, rows, x, y, w, max_h):
    nc = max(len(r) for r in rows)
    rows = [r + [""] * (nc - len(r)) for r in rows]
    nr = len(rows)
    weights = [max(4, min(30, max(text_w(r[j], 10) / 10 for r in rows))) for j in range(nc)]
    col_w = [w * wt / sum(weights) for wt in weights]

    def row_heights(font):
        hs = []
        for r in rows:
            lines = max(max(1, math.ceil(text_w(c, font) / max(col_w[j] * 72 - 14, 20))) for j, c in enumerate(r))
            hs.append(lines * font * 1.3 / 72 + 0.2)
        return hs

    fonts = (14, 13, 12, 11, 10) if nc < 6 else (12, 11, 10, 9)
    for font in fonts:
        hs = row_heights(font)
        if sum(hs) <= max_h:
            break
    shape = slide.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(sum(hs)))
    tbl = shape.table
    tbl.first_row = True
    tbl.horz_banding = False
    for j in range(nc):
        tbl.columns[j].width = Inches(col_w[j])
    # 证据等级色标：表头含"证据"时，单元格匹配 A/B/C 等级上色
    evidence_col = None
    for j, h in enumerate(rows[0]):
        if "证据" in h or "等级" in h:
            evidence_col = j
            break
    EVID_COLOR = {"A": RGBColor(0x2E, 0x7D, 0x32), "B": NAVY, "C": RGBColor(0x78, 0x84, 0x94)}

    for i, r in enumerate(rows):
        tbl.rows[i].height = Inches(hs[i])
        for j, val in enumerate(r):
            c = tbl.cell(i, j)
            c.margin_left = c.margin_right = Inches(0.1)
            c.margin_top = c.margin_bottom = Inches(0.05)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.fill.solid()
            c.fill.fore_color.rgb = NAVY if i == 0 else (WHITE if i % 2 else ROW_ALT)
            p = c.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if (j == 0 and i > 0) or (i == 0 and j == 0) else PP_ALIGN.CENTER
            run = p.add_run()
            run.text = val
            if i == 0:
                style_run(run, font, True, WHITE)
            elif j == 0:
                style_run(run, font, True, NAVY)
            elif evidence_col is not None and j == evidence_col:
                # 从单元格内容提取 A/B/C 等级首字母
                m = re.match(r"^([ABC])", val.strip())
                if m and m.group(1) in EVID_COLOR:
                    style_run(run, font, True, EVID_COLOR[m.group(1)])
                else:
                    style_run(run, font, False, INK)
            else:
                style_run(run, font, False, INK)
    return sum(hs)


def num(v):
    m = re.search(r"-?\d+(?:\.\d+)?", str(v).replace(",", ""))
    return float(m.group()) if m else None


def add_chart(slide, rows, ctype, title, x, y, w, h):
    cats = [r[0] for r in rows[1:]]
    names = [n or f"系列{k + 1}" for k, n in enumerate(rows[0][1:])]
    series = [[num(r[j]) for r in rows[1:]] for j in range(1, len(rows[0]))]
    if ctype == "bar":  # 横向条形图：让第一类别显示在最上方
        cats = cats[::-1]
        series = [v[::-1] for v in series]
    cd = CategoryChartData()
    cd.categories = cats
    for n_, v in zip(names, series):
        cd.add_series(n_, v)
    kind = {"bar": XL_CHART_TYPE.BAR_CLUSTERED, "column": XL_CHART_TYPE.COLUMN_CLUSTERED, "line": XL_CHART_TYPE.LINE_MARKERS}[ctype]
    ch = slide.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w), Inches(h), cd).chart
    ch.font.size = Pt(11)
    ch.font.color.rgb = GRAY
    ch.font.name = FONT
    if title:
        ch.has_title = True
        ch.chart_title.include_in_layout = False
        ch.chart_title.text_frame.text = title
        for r in ch.chart_title.text_frame.paragraphs[0].runs:
            style_run(r, 12, True, NAVY)
    else:
        ch.has_title = False
    ch.has_legend = len(names) > 1
    if ch.has_legend:
        ch.legend.position = XL_LEGEND_POSITION.BOTTOM
        ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(11)
    plot = ch.plots[0]
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.font.size = Pt(11)
    dl.font.bold = True
    dl.font.color.rgb = NAVY
    dl.position = XL_LABEL_POSITION.ABOVE if ctype == "line" else XL_LABEL_POSITION.OUTSIDE_END
    if ctype != "line":
        plot.gap_width = 55
        plot.vary_by_categories = False
    colors = [NAVY, CYAN, CORAL, GRAY]
    for i, ser in enumerate(plot.series):
        c = colors[i % 4]
        if ctype == "line":
            ser.format.line.color.rgb = c
            ser.format.line.width = Pt(2.5)
            ser.smooth = False
            ser.marker.format.fill.solid()
            ser.marker.format.fill.fore_color.rgb = c
            ser.marker.format.line.color.rgb = c
        else:
            ser.format.fill.solid()
            ser.format.fill.fore_color.rgb = c
    va, ca = ch.value_axis, ch.category_axis
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = LINE
    va.format.line.fill.background()
    va.tick_labels.font.size = Pt(10)
    allv = [v for s_ in series for v in s_ if v is not None]
    if allv and min(allv) >= 0:
        va.minimum_scale = 0
    ca.has_major_gridlines = False
    ca.format.line.color.rgb = LINE
    ca.tick_labels.font.size = Pt(11)


def add_images(slide, imgs, x, y, w, h):
    n = len(imgs)
    if n == 1:
        cols = 1
    elif n == 2:
        cols = 2 if w >= h * 1.3 else 1
    elif n <= 4:
        cols = 2
    else:
        cols = 3
    rows = math.ceil(n / cols)
    gap = 0.2
    cw = (w - gap * (cols - 1)) / cols
    ch = (h - gap * (rows - 1)) / rows
    for k, im in enumerate(imgs):
        r, c = divmod(k, cols)
        cx, cy = x + c * (cw + gap), y + r * (ch + gap)
        cap = (im.get("caption") or "").strip()
        cap_h = 0.5 if (cap and ch > 2.2) else 0
        bw, bh = cw, ch - cap_h
        scale = min(bw / im["w"], bh / im["h"])
        pw, ph_ = im["w"] * scale, im["h"] * scale
        px, py = cx + (bw - pw) / 2, cy + (bh - ph_) / 2
        slide.shapes.add_picture(io.BytesIO(im["bytes"]), Inches(px), Inches(py), Inches(pw), Inches(ph_))
        if cap_h:
            text = cap if len(cap) <= 130 else cap[:129] + "…"
            add_text(slide, cx, py + ph_ + 0.04, cw, cap_h - 0.04, text, 10, NAVY, align=PP_ALIGN.CENTER)


# ---------------- 各版式 ----------------
def L_bullets(slide, s, box, imgs):
    x, y, w, h = box
    lines = parse_bullets(s["bullets"])
    if imgs:
        lw = w * 0.46
        add_panel(slide, x, y, lw, h, s["sub_title"], lines)
        add_images(slide, imgs, x + lw + 0.4, y, w - lw - 0.4, h)
    else:
        add_panel(slide, x, y, w, h, s["sub_title"], lines, max_size=18)


def L_image(slide, s, box, imgs):
    x, y, w, h = box
    lines = parse_bullets(s["bullets"])
    if lines or s["sub_title"]:
        side = 3.75
        add_images(slide, imgs, x, y, w - side - 0.45, h)
        add_panel(slide, x + w - side, y + 0.1, side, h - 0.1, s["sub_title"] or "关键读数", lines, max_size=14)
    else:
        add_images(slide, imgs, x, y, w, h)


def L_chart(slide, s, box, imgs):
    x, y, w, h = box
    lines = parse_bullets(s["bullets"])
    rows = final_rows(s, "chart")
    side = bool(lines or s["sub_title"])
    lw = w * 0.62 if side else w
    add_chart(slide, rows, s["chart_type"], s["chart_title"], x, y, lw, h)
    if side:
        add_panel(slide, x + lw + 0.45, y + 0.3, w - lw - 0.45, h - 0.3, s["sub_title"] or "关键读数", lines, max_size=14)


def L_table(slide, s, box, imgs):
    x, y, w, h = box
    lines = parse_bullets(s["bullets"])
    rows = final_rows(s, "table")
    lw = w * 0.6 if imgs else w
    top = 0
    if lines or s["sub_title"]:
        pb = min(panel_need(s["sub_title"], lines, lw), h * 0.32)
        add_panel(slide, x, y, lw, pb, s["sub_title"], lines, max_size=15)
        top = pb + 0.2
    if rows:
        add_table(slide, rows, x, y + top, lw, h - top)
    if imgs:
        add_images(slide, imgs, x + lw + 0.4, y, w - lw - 0.4, h)


def L_stats(slide, s, box, imgs):
    x, y, w, h = box
    items = s["stats"][:4]
    n = len(items)
    gap = 0.5
    cw = (w - gap * (n - 1)) / n
    colors = [NAVY, CYAN, CORAL, NAVY]
    for i, d in enumerate(items):
        cx = x + i * (cw + gap)
        add_text(slide, cx, y, cw, 0.4, d["label"], 13, GRAY)
        size = next((z for z in (36, 32, 28, 24, 20, 18) if text_w(d["value"], z, True) <= cw * 72 - 16), 18)
        add_text(slide, cx, y + 0.42, cw, 0.8, d["value"], size, colors[i % 4], True, anchor=MSO_ANCHOR.MIDDLE)
        add_text(slide, cx, y + 1.25, cw, 0.6, d["sub"], 13, GRAY)
    rect(slide, x, y + 1.95, w, 0.015, LINE)
    lines = parse_bullets(s["bullets"])
    if (lines or s["sub_title"]) and h > 2.7:
        add_panel(slide, x, y + 2.2, w, h - 2.2, s["sub_title"], lines, max_size=16)


def L_flow(slide, s, box, imgs):
    x, y, w, h = box
    steps = s["steps"][:6]
    n = len(steps)
    # 判断哪些步骤是"决策点"：title 含 如果/若/当/是否/?/？ 或 text 含分支
    def is_decision(d):
        t = (d.get("title", "") + d.get("text", ""))
        return bool(re.search(r"如果|若|当.*时|是否|\?|？|≥|≤|<|>|提示|触发", t))

    ag = 0.55
    bw = (w - ag * (n - 1)) / n
    need = max(text_height([(d["text"], 0)], bw - 0.4, 13, indent=14) + 0.45 for d in steps)
    bh = min(max(1.5, need + 0.3), 2.6)
    for i, d in enumerate(steps):
        bx = x + i * (bw + ag)
        last = i == n - 1 and n >= 3
        decision = is_decision(d)
        # 决策点用菱形/六边形，终点用珊瑚色，普通步骤用卡片
        if decision:
            sh = slide.shapes.add_shape(MSO_SHAPE.DIAMOND, Inches(bx), Inches(y + 0.15), Inches(bw), Inches(bh - 0.1))
            sh.fill.solid()
            sh.fill.fore_color.rgb = RGBColor(0xFF, 0xF3, 0xE0)  # 浅橙黄
            sh.line.color.rgb = RGBColor(0xE7, 0x6B, 0x6B)
            sh.line.width = Pt(1.25)
            sh.shadow.inherit = False
            tf = sh.text_frame
            tf.word_wrap = True
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf.margin_left = tf.margin_right = Inches(0.08)
            tf.margin_top = tf.margin_bottom = Inches(0.04)
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            r = p.add_run()
            r.text = d["title"]
            style_run(r, 12, True, RGBColor(0xC0, 0x50, 0x20))
            if d.get("text"):
                p2 = tf.add_paragraph()
                p2.alignment = PP_ALIGN.CENTER
                r2 = p2.add_run()
                r2.text = d["text"]
                style_run(r2, 10, False, GRAY)
        else:
            rect(slide, bx, y + 0.1, bw, bh, CORAL_TINT if last else CARD)
            tf = textbox(slide, bx + 0.18, y + 0.2, bw - 0.36, bh - 0.2, MSO_ANCHOR.MIDDLE)
            t = d["title"] if re.match(r"^\d", d["title"]) else f"{i + 1}. {d['title']}"
            para(tf, 0, t, 13, True, CORAL if last else NAVY, after=4)
            para(tf, 1, d["text"], 13, last, CORAL if last else GRAY)
        if i < n - 1:
            add_text(slide, bx + bw, y + 0.1 + bh / 2 - 0.25, ag, 0.5, "→", 20, GRAY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    lines = parse_bullets(s["bullets"])
    top = bh + 0.55
    if (lines or s["sub_title"]) and h - top > 0.6:
        add_panel(slide, x, y + top, w, h - top, s["sub_title"], lines, max_size=15)


def L_cards(slide, s, box, imgs):
    x, y, w, h = box
    blocks = s["blocks"][:4]
    n = len(blocks)
    gap = 0.4
    cw = (w - gap * (n - 1)) / n
    inner = cw - 0.5
    cap = h * 0.7
    size = 13
    for size in (14, 13, 12, 11):
        need = max(0.6 + text_height([(l, 0) for l in b["lines"]], inner, size, indent=14, after=0.4) for b in blocks)
        if need + 0.3 <= cap:
            break
    ch = min(max(1.6, need + 0.3), cap)
    for i, b in enumerate(blocks):
        cx = x + i * (cw + gap)
        rect(slide, cx, y, cw, ch, CARD)
        tf = textbox(slide, cx + 0.25, y + 0.2, inner, ch - 0.3)
        para(tf, 0, b["heading"], 15, True, NAVY, after=8)
        for k, l in enumerate(b["lines"], start=1):
            para(tf, k, l.strip(), size, False, GRAY, after=size * 0.4)
    lines = parse_bullets(s["bullets"])
    top = ch + 0.4
    if (lines or s["sub_title"]) and h - top > 0.6:
        add_panel(slide, x, y + top, w, h - top, s["sub_title"], lines, max_size=15)


def L_two_col(slide, s, box, imgs):
    x, y, w, h = box
    blocks = s["blocks"][:3]
    n = len(blocks)
    gap = 0.9
    cw = (w - gap * (n - 1)) / n
    for i, b in enumerate(blocks):
        cx = x + i * (cw + gap)
        hs = next((z for z in (18, 16, 15, 14) if text_w(b["heading"], z) <= cw * 72 - 10), 14)
        color = CORAL if (s["contrast"] and i == 1) else NAVY
        add_text(slide, cx, y, cw, 0.5, b["heading"], hs, color, True)
        add_bullets(slide, cx, y + 0.65, cw, h - 0.65, parse_bullets(b["lines"]), max_size=15)
        if i < n - 1:
            rect(slide, cx + cw + gap / 2 - 0.01, y + 0.05, 0.015, min(h - 0.1, 3.6), LINE)


def L_numbered(slide, s, box, imgs):
    x, y, w, h = box
    items = s["items"][:6]
    n = len(items)
    rh = min(1.15, h / n)
    for i, t in enumerate(items):
        ry = y + i * rh
        add_text(slide, x, ry, 1.1, rh, f"{i + 1:02d}", 30, CYAN, True, anchor=MSO_ANCHOR.MIDDLE)
        size = next((z for z in (16, 15, 14, 13, 12) if text_height([(t, 0)], w - 1.35, z, indent=14, after=0) <= rh - 0.1), 12)
        add_text(slide, x + 1.3, ry, w - 1.3, rh, t, size, NAVY, True, anchor=MSO_ANCHOR.MIDDLE)


def L_grid(slide, s, box, imgs):
    """2×2 编号块：blocks 3-4 项，每项 heading + lines。"""
    x, y, w, h = box
    blocks = s["blocks"][:4]
    n = len(blocks)
    if n == 1:
        cols, rows = 1, 1
    elif n == 2:
        cols, rows = 2, 1
    elif n == 3:
        cols, rows = 3, 1
    else:
        cols, rows = 2, 2
    gap = 0.45
    cw = (w - gap * (cols - 1)) / cols
    ch = (h - gap * (rows - 1)) / rows
    for i, b in enumerate(blocks):
        r, c = divmod(i, cols)
        cx = x + c * (cw + gap)
        cy = y + r * (ch + gap)
        add_text(slide, cx, cy, 0.9, 0.5, f"{i + 1:02d}", 26, CYAN, True, anchor=MSO_ANCHOR.TOP)
        add_text(slide, cx + 1.0, cy, cw - 1.0, 0.4, b["heading"], 15, NAVY, True)
        lines = parse_bullets(b["lines"])
        if lines:
            add_bullets(slide, cx + 1.0, cy + 0.5, cw - 1.0, ch - 0.5, lines, max_size=13)


def L_pillars(slide, s, box, imgs):
    """2-4 个并列编号柱：blocks 每项 heading + lines，顶部大编号。"""
    x, y, w, h = box
    blocks = s["blocks"][:4]
    n = len(blocks) or 1
    gap = 0.5
    cw = (w - gap * (n - 1)) / n
    for i, b in enumerate(blocks):
        cx = x + i * (cw + gap)
        add_text(slide, cx, y, cw, 0.6, f"{i + 1:02d}", 34, NAVY, True)
        rect(slide, cx, y + 0.7, 0.6, 0.04, CORAL)
        add_text(slide, cx, y + 0.85, cw, 0.5, b["heading"], 15, NAVY, True)
        lines = parse_bullets(b["lines"])
        if lines:
            add_bullets(slide, cx, y + 1.4, cw, h - 1.4, lines, max_size=13)


def L_chains(slide, s, box, imgs):
    """多条因果链：每条链 heading + 步骤 bullets（3-5 步，步骤间用 → 连接）。"""
    x, y, w, h = box
    blocks = s["blocks"][:4]
    n = len(blocks)
    rh = h / max(n, 1)
    for i, b in enumerate(blocks):
        ry = y + i * rh
        add_text(slide, x, ry, 1.6, 0.35, b["heading"], 13, NAVY, True)
        steps = [l.strip().lstrip("-–•·*").strip() for l in b["lines"] if l.strip()]
        if not steps:
            continue
        # 横向排步骤，用 → 连接
        seg_w = (w - 1.8 - 0.3 * (len(steps) - 1)) / len(steps)
        sx = x + 1.8
        for k, step in enumerate(steps):
            size = next((z for z in (12, 11, 10) if text_w(step, z) <= seg_w * 72 - 6), 10)
            color = CORAL if k == len(steps) - 1 else GRAY
            add_text(slide, sx + k * (seg_w + 0.3), ry, seg_w, 0.5, step, size, color,
                     bold=(k == len(steps) - 1), anchor=MSO_ANCHOR.MIDDLE)
            if k < len(steps) - 1:
                add_text(slide, sx + k * (seg_w + 0.3) + seg_w, ry, 0.3, 0.5, "→", 14, GRAY,
                         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        if i < n - 1:
            rect(slide, x, ry + rh - 0.02, w, 0.01, LINE)


def L_abstract(slide, s, box, imgs):
    """摘要页：左栏 blocks（带标签段落），右栏 items 关键词。"""
    x, y, w, h = box
    blocks = s["blocks"]
    items = s["items"]
    side = 2.6 if items else 0
    lw = w - side - 0.4 if items else w
    cy = y
    for b in blocks:
        if b["heading"]:
            add_text(slide, x, cy, lw, 0.35, b["heading"], 14, NAVY, True)
            cy += 0.4
        lines = b["lines"]
        if lines:
            tf = textbox(slide, x, cy, lw, 0.3 * len(lines) + 0.2)
            for k, l in enumerate(lines):
                para(tf, k, l.strip(), 13, False, GRAY, after=4)
            cy += 0.32 * len(lines) + 0.15
    if items:
        add_text(slide, x + lw + 0.4, y, side, 0.35, "关键词", 13, NAVY, True)
        rect(slide, x + lw + 0.4, y + 0.4, 0.8, 0.03, CORAL)
        tf = textbox(slide, x + lw + 0.4, y + 0.55, side, h - 0.55)
        for k, it in enumerate(items):
            para(tf, k, it.strip(), 12, False, GRAY, after=6)


LAYOUT_FUNCS = dict(
    bullets=L_bullets, image=L_image, chart=L_chart, table=L_table, stats=L_stats,
    flow=L_flow, cards=L_cards, two_col=L_two_col, numbered=L_numbered,
    grid=L_grid, pillars=L_pillars, chains=L_chains, abstract=L_abstract,
)


# ---------------- 页面框架 ----------------
class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(SW), Inches(SH)
        bg = self.prs.slide_master.background.fill
        bg.solid()
        bg.fore_color.rgb = BG
        self.page = 0

    def blank(self):
        self.page += 1
        return self.prs.slides.add_slide(self.prs.slide_layouts[5])

    def frame(self, title, label="", title_size=None):
        slide = self.blank()
        if label:
            add_text(slide, X, 0.5, 8, 0.33, label, 12, GRAY)
        t = title or " "
        size = title_size or next((z for z in (28, 26, 24, 22, 20, 18) if text_w(t, z) <= W * 72 - 14), 18)
        set_title(slide.shapes.title, t, size, X, 0.86, W, 0.78, MSO_ANCHOR.MIDDLE)
        rect(slide, X, 1.72, 1.4, 0.045, CORAL)
        add_text(slide, SW - 1.6, 7.05, 0.7, 0.3, str(self.page), 10, GRAY, align=PP_ALIGN.RIGHT)
        return slide

    def cover(self, c):
        slide = self.blank()
        rect(slide, 0, 0, 4.44, SH, NAVY)
        rect(slide, 0, 6.39, 4.44, 1.11, CYAN)
        if c["kicker"]:
            add_text(slide, 5.0, 1.2, 7.5, 0.9, c["kicker"], 14, CYAN, anchor=MSO_ANCHOR.BOTTOM)
        title = c["title"] or "文献汇报"
        tw = 7.5 * 72 - 14
        size = 22
        for z in (32, 30, 28, 26, 24, 22):
            n = sum(max(1, math.ceil(text_w(l, z) / tw)) for l in title.split("\n"))
            if n * z * 1.3 / 72 <= 2.2:
                size = z
                break
        set_title(slide.shapes.title, title, size, 5.0, 2.25, 7.5, 2.25, MSO_ANCHOR.TOP)
        rect(slide, 5.0, 4.7, 2.22, 0.045, CORAL)
        if c["meta"]:
            add_text(slide, 5.0, 4.95, 7.5, 1.1, c["meta"], 14, GRAY)
        if c["foot"]:
            add_text(slide, 5.0, 6.3, 7.5, 0.6, c["foot"], 12, GRAY)

    def agenda(self, secs):
        slide = self.frame("目录", title_size=28)
        n = len(secs)
        rows = math.ceil(n / 2)
        rh = min(1.53, 4.6 / rows)
        for k, sc in enumerate(secs):
            r, c = divmod(k, 2)
            x = 1.1 + c * 6.1
            y = 2.1 + r * rh
            add_text(slide, x, y, 0.9, 0.6, f"{k + 1:02d}", 24, CYAN, True)
            add_text(slide, x + 1.1, y, 4.6, 0.45, sc["title"], 18, NAVY, True)
            if sc["subtitle"]:
                add_text(slide, x + 1.1, y + 0.45, 4.6, 0.6, sc["subtitle"], 12, GRAY)
            if c == 1 or k == n - 1:
                if r < rows - 1:
                    rect(slide, X, y + rh - 0.35, W, 0.012, LINE)

    def content(self, o, s, images):
        slide = self.frame(s["title"], sec_label(o, s["sec"]))
        imgs = [images[i] for i in s["image_ids"] if i in images] if s["layout"] in ("image", "table", "bullets") else []
        call = s["callout"].strip()
        y1 = Y1
        if call:
            ch = callout_h(call, W)
            add_callout(slide, call, X, Y1 - ch, W, ch)
            y1 = Y1 - ch - 0.3
        LAYOUT_FUNCS[s["layout"]](slide, s, (X, Y0, W, y1 - Y0), imgs)
        if s["notes"].strip():
            slide.notes_slide.notes_text_frame.text = s["notes"]

    def references(self, refs):
        slide = self.frame("参考文献", title_size=28)
        add_text(slide, X, Y0, W, 0.45, "主要文献依据", 15, NAVY, True)
        lines = [f"{i + 1}. {r}" for i, r in enumerate(refs)]
        size = next((z for z in (14, 13, 12, 11, 10) if text_height([(l, 0) for l in lines], W, z, indent=14, after=0.6) <= Y1 - Y0 - 0.7), 10)
        add_plain(slide, X, Y0 + 0.6, W, Y1 - Y0 - 0.6, lines, GRAY, size)

    def save(self):
        buf = io.BytesIO()
        self.prs.save(buf)
        return buf.getvalue()


def build_pptx(o, images):
    d = Deck()
    d.cover(o["cover"])
    if o["show_agenda"] and len(o["sections"]) >= 2:
        d.agenda(o["sections"])
    for s in o["slides"]:
        d.content(o, s, images)
    if o["show_refs"] and o["references"]:
        d.references(o["references"])
    return d.save()


# ============================================================================
# 五、GitHub 保存
# ============================================================================
def github_put(token, repo, branch, path, content: bytes, message):
    url = f"https://api.github.com/repos/{repo}/contents/{urllib.parse.quote(path)}"
    headers = {
        "Authorization": f"Bearer {token}",
        "Accept": "application/vnd.github+json",
        "X-GitHub-Api-Version": "2022-11-28",
    }
    r = requests.get(url, headers=headers, params={"ref": branch}, timeout=30)
    sha = r.json().get("sha") if r.status_code == 200 else None
    payload = {"message": message, "content": base64.b64encode(content).decode(), "branch": branch}
    if sha:
        payload["sha"] = sha
    r = requests.put(url, headers=headers, json=payload, timeout=120)
    if r.status_code not in (200, 201):
        try:
            msg = r.json().get("message", r.text)
        except Exception:
            msg = r.text
        raise RuntimeError(f"GitHub 返回 {r.status_code}：{msg[:200]}")
    return r.json()["content"]["html_url"]


def outline_to_json(o):
    drop = ("uid", "table_cur", "chart_cur", "tv", "cv")
    slides = []
    for s in o["slides"]:
        d = {k: v for k, v in s.items() if k not in drop}
        d["table"], d["chart"] = final_rows(s, "table"), final_rows(s, "chart")
        slides.append(d)
    return json.dumps(
        dict(cover=o["cover"], sections=o["sections"], slides=slides, references=o["references"]),
        ensure_ascii=False, indent=2,
    )


# ============================================================================
# 六、编辑器
# ============================================================================
def commit_tables():
    """把表格/图表数据的编辑结果固化为新基准，并更换 key，避免 data_editor 状态叠加。"""
    if not st.session_state.outline:
        return
    for s in st.session_state.outline["slides"]:
        for field, cur, ver in (("table", "table_cur", "tv"), ("chart", "chart_cur", "cv")):
            if s.get(cur) is not None:
                s[field], s[cur] = s[cur], None
                s[ver] += 1


def cb_move(i, d):
    commit_tables()
    sl = st.session_state.outline["slides"]
    j = i + d
    if 0 <= j < len(sl):
        sl[i], sl[j] = sl[j], sl[i]
        st.session_state.sel = j


def cb_delete(i):
    commit_tables()
    sl = st.session_state.outline["slides"]
    if sl:
        sl.pop(i)
        st.session_state.sel = max(0, min(i, len(sl) - 1))


def cb_add(i):
    commit_tables()
    sl = st.session_state.outline["slides"]
    sec = sl[i]["sec"] if 0 <= i < len(sl) else 0
    sl.insert(i + 1, new_slide(title="新页面", bullets=["要点 1"], sec=sec))
    st.session_state.sel = i + 1


def get_thumb(im):
    if "thumb" not in im:
        t = Image.open(io.BytesIO(im["bytes"]))
        t.thumbnail((320, 320))
        buf = io.BytesIO()
        t.convert("RGB").save(buf, "JPEG", quality=80)
        im["thumb"] = buf.getvalue()
    return im["thumb"]


def img_label(images, i):
    im = images[i]
    cap = (im.get("caption") or "")[:24]
    pg = f"第{im['page']}页 " if im.get("page") else ""
    return f"图{i}（{pg}{cap}）"


def grid_editor(s, field, default_rows, note):
    cur, ver = ("table_cur", "tv") if field == "table" else ("chart_cur", "cv")
    uid = s["uid"]
    if not s[field]:
        s[field] = default_rows
    base = s[field]
    nc = len(base[0])
    st.caption(note)
    new_nc = st.number_input("列数", 1, 10, nc, key=f"nc_{field}_{uid}")
    if new_nc != nc:
        src = s[cur] or s[field]
        s[field] = [(list(r) + [""] * new_nc)[:new_nc] for r in src]
        s[cur] = None
        s[ver] += 1
        st.rerun()
    df = pd.DataFrame(base, columns=[f"列{j + 1}" for j in range(nc)]).astype(str)
    edited = st.data_editor(df, num_rows="dynamic", width="stretch", key=f"ed_{field}_{uid}_{nc}_{s[ver]}")
    s[cur] = edited.fillna("").astype(str).values.tolist()


def render_slide_editor(o, s, images):
    uid = s["uid"]
    c1, c2 = st.columns([2, 1])
    with c1:
        s["title"] = st.text_input("页面标题（建议写成结论句）", s["title"], key=f"t_{uid}")
    with c2:
        opts = list(range(len(o["sections"]) + 1))
        s["sec"] = st.selectbox(
            "所属章节", opts, index=min(s["sec"], len(opts) - 1), key=f"sec_{uid}",
            format_func=lambda k: "（无）" if k == 0 else sec_label(o, k),
        )
    keys = list(LAYOUTS)
    s["layout"] = st.selectbox(
        "版式", keys, index=keys.index(s["layout"]), key=f"ly_{uid}", format_func=lambda k: LAYOUTS[k]
    )
    lay = s["layout"]

    if lay in ("two_col", "cards"):
        st.caption("每个栏目以“## 标题”开头，下面每行一条内容（二栏版式中，行首空两格表示二级要点）。")
        s["blocks"] = parse_blocks(st.text_area("栏目内容", blocks_to_text(s["blocks"]), height=260, key=f"bl_{uid}"))
        if lay == "two_col":
            s["contrast"] = st.checkbox("第二栏标题使用强调色（红）", s["contrast"], key=f"ct_{uid}")
    elif lay == "stats":
        st.caption("每行一个指标：标签 | 数值 | 说明（建议 2-4 行）。")
        s["stats"] = parse_stats(st.text_area("大数字指标", stats_to_text(s["stats"]), height=140, key=f"st_{uid}"))
    elif lay == "flow":
        st.caption("每行一个步骤：标题 | 说明（建议 3-5 行，最后一步会用红色强调）。")
        s["steps"] = parse_steps(st.text_area("流程步骤", steps_to_text(s["steps"]), height=160, key=f"sp_{uid}"))
    elif lay == "numbered":
        st.caption("每行一条结论。")
        s["items"] = str_lines(st.text_area("编号结论", "\n".join(s["items"]), height=200, key=f"it_{uid}"))

    if lay not in ("two_col", "numbered"):
        if lay in ("image", "chart"):
            s["sub_title"] = st.text_input("读数面板标题", s["sub_title"] or "关键读数", key=f"sb_{uid}")
        else:
            s["sub_title"] = st.text_input("小标题（可选）", s["sub_title"], key=f"sb_{uid}")
        txt = st.text_area(
            "要点（每行一条；行首空两格 = 二级要点）", "\n".join(s["bullets"]), height=170, key=f"b_{uid}",
        )
        s["bullets"] = [l.rstrip() for l in txt.split("\n") if l.strip()]

    if lay in ("image", "table", "bullets") and images:
        st.markdown("**本页使用的原文图片**（勾选缩略图下方的复选框）")
        opts = sorted(images)
        chosen = []
        for row_start in range(0, len(opts), 5):
            cols = st.columns(5)
            for col, iid in zip(cols, opts[row_start : row_start + 5]):
                with col:
                    st.image(get_thumb(images[iid]), width="stretch")
                    if st.checkbox(img_label(images, iid), iid in s["image_ids"], key=f"pick_{uid}_{iid}"):
                        chosen.append(iid)
        s["image_ids"] = chosen
        if chosen:
            st.markdown("**已选图片的图注**（显示在图片下方，可修改或清空）")
            cols = st.columns(min(3, len(chosen)))
            for k, iid in enumerate(chosen):
                with cols[k % len(cols)]:
                    images[iid]["caption"] = st.text_input(f"图{iid} 图注", images[iid].get("caption", ""), key=f"cap_{iid}")

    if lay == "table":
        grid_editor(s, "table", [["指标", "组 A", "组 B"], ["", "", ""], ["", "", ""]], "第一行为表头；可直接增删行、修改单元格。")
    if lay == "chart":
        c1, c2 = st.columns([1, 2])
        types = ["bar", "column", "line"]
        names = {"bar": "横向条形图", "column": "柱状图", "line": "折线图"}
        s["chart_type"] = c1.selectbox("图表类型", types, index=types.index(s["chart_type"]), format_func=lambda k: names[k], key=f"ctp_{uid}")
        s["chart_title"] = c2.text_input("图表标题（含单位）", s["chart_title"], key=f"cti_{uid}")
        grid_editor(s, "chart", [["类别", "系列1"], ["A", "10"], ["B", "20"], ["C", "30"]],
                    "第一行是表头（第一列为类别，其余列为数据系列）；数值列请填写数字。生成后是 PowerPoint 原生图表，可继续编辑数据。")

    s["callout"] = st.text_area("页面底部「关键发现」框（可选，留空则不显示）", s["callout"], height=80, key=f"co_{uid}")
    s["notes"] = st.text_area("讲稿（写入 PPT 备注）", s["notes"], height=160, key=f"n_{uid}")


# ============================================================================
# 七、界面
# ============================================================================
def main():
    for _k, _v in dict(outline=None, images={}, ppt=None, source="", sel=0, added=set()).items():
        st.session_state.setdefault(_k, _v)

    st.title("🩺 医学文献 → 学术汇报 PPT")
    st.caption("上传文献 → AI 提取大纲与数据 → 网页内逐页编辑 → 生成原生可编辑 PPT（封面、目录、结论式标题、数据卡片、图表、表格、讲稿）")

    with st.sidebar:
        st.header("⚙️ 模型设置")
        try:
            secret_key = st.secrets.get("OPENROUTER_API_KEY", "")
        except Exception:
            secret_key = ""
        api_key = st.text_input("OpenRouter API Key", type="password", placeholder="sk-or-v1-...") or secret_key
        free_live = fetch_free_models()
        free_pool = free_live or FREE_MODELS
        if DEFAULT_MODEL not in free_pool:
            free_pool = [DEFAULT_MODEL] + free_pool
        if "openrouter/free" not in free_pool:
            free_pool = free_pool + ["openrouter/free"]
        st.caption(f"已从 OpenRouter 实时获取 {len(free_live)} 个免费模型（上下文 ≥32k）" if free_live else "未能联网获取模型列表，使用内置备用列表")
        _model_choices = free_pool + PAID_MODELS
        _default_idx = _model_choices.index(DEFAULT_MODEL) if DEFAULT_MODEL in _model_choices else 0
        model_name = st.selectbox("选择模型", _model_choices, index=_default_idx)
        custom_model = st.text_input("自定义模型 ID（可选，优先使用）", placeholder="如 anthropic/claude-sonnet-4.5")
        auto_fallback = st.checkbox("失败时自动尝试其他免费模型", True)
        max_chars = st.slider("送入模型的最大字符数", 10000, 120000, 45000, 5000,
                              help="免费模型上下文较短，长文献请调小或换用大上下文模型。")
        st.markdown("---")
        st.header("☁️ GitHub 保存")
        try:
            gh_token_secret = st.secrets.get("GITHUB_TOKEN", "")
            gh_repo_secret = st.secrets.get("GITHUB_REPO", "")
        except Exception:
            gh_token_secret, gh_repo_secret = "", ""
        gh_token = st.text_input("GitHub Token", type="password", placeholder="github_pat_...",
                                 help="需要对目标仓库有 Contents: Read and write 权限。") or gh_token_secret
        gh_repo = st.text_input("仓库（owner/repo）", value=gh_repo_secret, placeholder="your-name/ppt-output")
        gh_branch = st.text_input("分支", value="main")
        gh_folder = st.text_input("保存目录", value="ppt_output")
        gh_with_json = st.checkbox("同时保存大纲 JSON", True)
        st.markdown("---")
        st.markdown("提示：扫描版 PDF 无文字层，需先做 OCR；矢量绘制的图表无法被提取，可在编辑阶段手动上传截图。")

    st.subheader("① 上传文献并生成大纲")
    uploaded = st.file_uploader("支持 PDF / Word(.docx) / Markdown / TXT", type=["pdf", "docx", "md", "txt"])
    extra = st.text_area("补充要求（可选）", placeholder="例如：用于科室组会汇报，重点讲方法学和主要结果，不讲背景细节。")
    c1, c2 = st.columns(2)
    n_slides = c1.slider("内容页数（不含封面、目录、参考文献）", 6, 30, 20)
    c1.caption(f"预计总页数约 {n_slides + 3} 页（内容页 + 封面 + 目录 + 参考文献）")
    lang = c2.selectbox("PPT 语言", ["中文", "English"])

    if st.button("🚀 分析文献并生成大纲", type="primary"):
        if not api_key:
            st.error("请先在左侧填写 OpenRouter API Key。")
        elif not uploaded:
            st.error("请先上传文献。")
        else:
            try:
                with st.status("正在处理…", expanded=True) as status:
                    st.write("解析文献（文本 / 图片 / 表格）…")
                    text, imgs = parse_file(uploaded)
                    text = strip_references(text).strip()
                    if len(text) < 200:
                        raise RuntimeError("几乎没有提取到文字，可能是扫描版 PDF，请先 OCR。")
                    images = {im["id"]: im for im in imgs}
                    st.write(f"提取到 {len(text)} 字符、{len(images)} 张图片。")
                    if len(text) > max_chars:
                        half = max_chars // 2
                        text = text[:half] + "\n\n[……中间部分因长度限制省略……]\n\n" + text[-half:]

                    img_desc = "\n".join(
                        f"图片{i}: " + (f"第{m['page']}页; " if m["page"] else "") + f"{m['w']}x{m['h']}; 原文图注: {m['caption'] or '无'}"
                        for i, m in images.items()
                    ) or "（无可用图片）"
                    lang_req = {"中文": "全部使用中文（专业缩写保留英文）", "English": "全部使用英文"}[lang]

                    # 加载医学术语对照表，挑出文献中实际出现的术语附给模型
                    st.write("加载医学术语对照表…")
                    gloss = load_glossary()
                    gloss_block = pick_glossary_for_text(text, gloss, limit=120)
                    if gloss_block:
                        st.write(f"  · 命中 {gloss_block.count(chr(10)) - 1} 个文献中出现的术语")
                    gloss_section = f"\n\n【医学术语标准译法】\n{gloss_block}\n" if gloss_block else ""

                    user_prompt = (
                        f"请为下面这篇医学文献生成完整的学术汇报内容（封面、章节目录、内容页、参考文献）。\n"
                        f"- 内容页约 {n_slides} 页（不含封面、目录和参考文献）\n- 语言：{lang_req}\n"
                        f"- 每个结论性要点都必须附带文献中的具体数据，数值逐一核对原文；版式要多样\n"
                        f"- 英文医学术语翻译成中文时，必须使用下面【医学术语标准译法】表中的译法，不要自创译法；表中未列出的再按医学惯例翻译\n"
                        f"- 额外要求：{extra.strip() or '无'}\n\n【图片清单】\n{img_desc}\n{gloss_section}\n【文献正文】\n{text}"
                    )
                    first = custom_model.strip() or model_name
                    models = [first] + ([m for m in free_pool[:6] if m != first] if auto_fallback else [])
                    st.write(f"调用模型 `{first}` 生成大纲…")
                    data, used = call_llm(api_key, models, user_prompt, lambda m, e: st.write(f"⚠️ `{m}`：{str(e)[:160]}"))
                    outline = normalize_outline(data, images)
                    if not outline["slides"]:
                        raise RuntimeError("模型返回的大纲为空，请重试或换模型。")
                    st.session_state.update(outline=outline, images=images, ppt=None, sel=0, source=uploaded.name, added=set())
                    status.update(label=f"大纲已生成（模型：{used}）", state="complete")
            except Exception as e:
                st.error(f"失败：{e}")

    outline = st.session_state.outline
    if outline:
        images = st.session_state.images
        st.markdown("---")
        st.subheader("② 编辑内容")

        with st.expander(f"📷 文献中可用的图片（{len(images)} 张）/ 手动补充图片"):
            extra_imgs = st.file_uploader("补充上传图片（如矢量图表的截图）", type=["png", "jpg", "jpeg"], accept_multiple_files=True, key="extra_imgs")
            for f in extra_imgs or []:
                sig = (f.name, f.size)
                if sig in st.session_state.added:
                    continue
                n = norm_image(f.getvalue())
                if n:
                    nid = max(images, default=0) + 1
                    images[nid] = dict(id=nid, page=None, caption=f.name.rsplit(".", 1)[0], bytes=n[0], w=n[1], h=n[2])
                    st.session_state.added.add(sig)
            if images:
                cols = st.columns(5)
                for k, iid in enumerate(sorted(images)):
                    with cols[k % 5]:
                        st.image(get_thumb(images[iid]), width="stretch")
                        st.caption(img_label(images, iid))
            else:
                st.info("未从文献中提取到图片。")

        ou = outline["uid"]
        with st.expander("🎬 封面"):
            cv = outline["cover"]
            cv["kicker"] = st.text_input("顶部小字（期刊 · 年份）", cv["kicker"], key=f"cv_k_{ou}")
            cv["title"] = st.text_area("标题（可换行）", cv["title"], height=100, key=f"cv_t_{ou}")
            cv["meta"] = st.text_area("作者 / 单位（可多行）", cv["meta"], height=80, key=f"cv_m_{ou}")
            cv["foot"] = st.text_input("底部一句话（研究样本/设计概括）", cv["foot"], key=f"cv_f_{ou}")

        with st.expander("📑 目录章节与参考文献"):
            outline["show_agenda"] = st.checkbox("生成目录页", outline["show_agenda"], key=f"sa_{ou}")
            sec_txt = st.text_area(
                "章节（每行一个：标题 | 副标题）",
                "\n".join(f"{x['title']} | {x['subtitle']}" for x in outline["sections"]), height=140, key=f"se_{ou}",
            )
            outline["sections"] = [
                dict(title=p[0], subtitle=p[1]) for p in ((split_bar(l) + [""])[:2] for l in sec_txt.split("\n") if l.strip())
            ]
            outline["show_refs"] = st.checkbox("生成参考文献页", outline["show_refs"], key=f"sr_{ou}")
            ref_txt = st.text_area("参考文献（每行一条）", "\n".join(outline["references"]), height=100, key=f"rf_{ou}")
            outline["references"] = str_lines(ref_txt)

        sl = outline["slides"]
        n = len(sl)
        with st.expander(f"📋 内容总览（共 {n} 页内容页）"):
            for k, s in enumerate(sl, start=1):
                lab = sec_label(outline, s["sec"])
                st.markdown(f"**{k}. {s['title']}**　`{LAYOUTS[s['layout']]}`" + (f"　{lab}" if lab else ""))

        if n:
            if st.session_state.sel >= n:
                st.session_state.sel = n - 1
            idx = st.selectbox(
                "选择要编辑的页面", range(n), key="sel", on_change=commit_tables,
                format_func=lambda i: f"第 {i + 1} 页 · {sl[i]['title'] or '(无标题)'}",
            )
            b1, b2, b3, b4, _ = st.columns([1, 1, 1, 1.4, 3])
            b1.button("⬆ 上移", on_click=cb_move, args=(idx, -1), disabled=idx == 0)
            b2.button("⬇ 下移", on_click=cb_move, args=(idx, 1), disabled=idx == n - 1)
            b3.button("🗑 删除", on_click=cb_delete, args=(idx,))
            b4.button("➕ 在其后新增", on_click=cb_add, args=(idx,))
            with st.container(border=True):
                render_slide_editor(outline, sl[idx], images)
        else:
            st.warning("没有页面了。")
            st.button("➕ 新增页面", on_click=cb_add, args=(-1,))

        st.markdown("---")
        st.subheader("③ 生成 PPT")
        if st.button("📊 生成 PowerPoint", type="primary"):
            try:
                st.session_state.ppt = build_pptx(outline, images)
            except Exception as e:
                st.error(f"生成失败：{e}")
        if st.session_state.ppt:
            st.success("生成成功！文字、表格、图表、图片均为原生对象，可在 PowerPoint / WPS 中继续编辑，讲稿在备注栏。")
            fname = re.sub(r'[\\/:*?"<>|\n]', "_", outline["cover"]["title"] or "presentation")[:40] + ".pptx"
            st.download_button(
                "📥 下载 .pptx", st.session_state.ppt, file_name=fname,
                mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
            )
            if st.button("☁️ 保存到 GitHub"):
                if not (gh_token and gh_repo.strip()):
                    st.error("请先在左侧填写 GitHub Token 和仓库（owner/repo）。")
                else:
                    try:
                        stamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
                        base = f"{stamp}_{fname[:-5]}"
                        folder = gh_folder.strip().strip("/")
                        prefix = f"{folder}/" if folder else ""
                        repo, branch = gh_repo.strip(), (gh_branch.strip() or "main")
                        with st.spinner("正在上传到 GitHub…"):
                            url = github_put(gh_token, repo, branch, f"{prefix}{base}.pptx",
                                             st.session_state.ppt, f"Add presentation: {base}")
                            if gh_with_json:
                                github_put(gh_token, repo, branch, f"{prefix}{base}.outline.json",
                                           outline_to_json(outline).encode("utf-8"), f"Add outline: {base}")
                        st.success(f"已保存到 GitHub：[{url}]({url})")
                    except Exception as e:
                        st.error(f"保存失败：{e}")


if __name__ == "__main__":
    main()
